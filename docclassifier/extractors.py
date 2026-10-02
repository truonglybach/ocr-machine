from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Protocol

from .models import ExtractedContent

MAX_CHARS = 20_000  # enough signal for classification, bounded memory


class ExtractionError(Exception):
    """Raised when a file cannot be read."""


class TextExtractor(Protocol):
    extensions: tuple[str, ...]

    def extract(self, path: Path) -> ExtractedContent: ...


def _as_date(value: datetime | date | None) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value


class DocxExtractor:
    extensions = (".docx",)

    def extract(self, path: Path) -> ExtractedContent:
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(c.text for c in row.cells if c.text.strip())
        props = doc.core_properties
        title = (props.title or "").strip() or (parts[0] if parts else None)
        return ExtractedContent("\n".join(parts)[:MAX_CHARS], title, _as_date(props.created))


class XlsxExtractor:
    extensions = (".xlsx", ".xlsm")

    def extract(self, path: Path) -> ExtractedContent:
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        try:
            parts: list[str] = []
            size = 0
            for ws in wb.worksheets:
                parts.append(ws.title)
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        line = " ".join(cells)
                        parts.append(line)
                        size += len(line)
                    if size >= MAX_CHARS:
                        break
                if size >= MAX_CHARS:
                    break
            title = (wb.properties.title or "").strip() or (wb.worksheets[0].title if wb.worksheets else None)
            return ExtractedContent("\n".join(parts)[:MAX_CHARS], title, _as_date(wb.properties.created))
        finally:
            wb.close()


class PlainTextExtractor:
    extensions = (".txt", ".md", ".csv")

    def extract(self, path: Path) -> ExtractedContent:
        text = path.read_text(encoding="utf-8", errors="replace")[:MAX_CHARS]
        first = next((ln.strip() for ln in text.splitlines() if ln.strip()), None)
        return ExtractedContent(text, first)


class ExtractorRegistry:
    """Maps file extensions to extractors; register more to support new types."""

    def __init__(self, extractors: list[TextExtractor] | None = None):
        self._by_ext: dict[str, TextExtractor] = {}
        for extractor in extractors or []:
            self.register(extractor)

    def register(self, extractor: TextExtractor) -> None:
        for ext in extractor.extensions:
            self._by_ext[ext.lower()] = extractor

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in self._by_ext

    def extract(self, path: Path) -> ExtractedContent:
        extractor = self._by_ext.get(path.suffix.lower())
        if extractor is None:
            raise ExtractionError(f"Unsupported file type: {path.suffix}")
        try:
            return extractor.extract(path)
        except Exception as exc:  # corrupt/locked/encrypted files must not abort a batch
            raise ExtractionError(f"{path.name}: {exc}") from exc


def default_registry() -> ExtractorRegistry:
    return ExtractorRegistry([DocxExtractor(), XlsxExtractor(), PlainTextExtractor()])
