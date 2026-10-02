from datetime import date

import pytest
from docx import Document
from openpyxl import Workbook

from docclassifier import KeywordClassifier, NamingConvention, RenamePipeline, default_registry
from docclassifier import RenameJournal, SkipReason


def make_docx(path, *lines):
    d = Document()
    for ln in lines:
        d.add_paragraph(ln)
    d.save(path)


def make_xlsx(path, title, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    wb.save(path)


@pytest.fixture
def tree(tmp_path):
    (tmp_path / "sub" / "deep").mkdir(parents=True)
    make_docx(tmp_path / "scan001.docx", "Acme Invoice", "Invoice number 42", "Amount due: $100", "Bill to: Bob")
    make_xlsx(tmp_path / "sub" / "book1.xlsx", "Q3 Budget", [["Revenue", 10], ["Expenses", 5], ["Forecast", 7]])
    make_docx(tmp_path / "sub" / "deep" / "x.docx", "Zebra", "nothing special here")
    (tmp_path / "sub" / "deep" / "bad.docx").write_text("not a real docx")
    (tmp_path / "image.png").write_bytes(b"\x89PNG")
    return tmp_path


def pipeline(template="{category}_{title}"):
    return RenamePipeline(default_registry(), KeywordClassifier(), NamingConvention(template))


def test_plan_recurses_classifies_and_skips(tree):
    plan = pipeline().plan(tree)
    names = {a.source.name: a.target.name for a in plan.actions}
    assert names["scan001.docx"] == "invoice_acme-invoice.docx"
    assert names["book1.xlsx"] == "budget_q3-budget.xlsx"
    assert names["x.docx"] == "other_zebra.docx"
    reasons = {s.path.name: s.reason for s in plan.skipped}
    assert reasons["bad.docx"] is SkipReason.UNREADABLE
    assert reasons["image.png"] is SkipReason.UNSUPPORTED


def test_dry_run_changes_nothing_and_apply_then_undo(tree):
    before = sorted(p.name for p in tree.rglob("*"))
    p = pipeline()
    plan = p.plan(tree)
    assert sorted(x.name for x in tree.rglob("*")) == before
    journal = RenameJournal(tree / "log.jsonl")
    result = p.apply(plan, journal)
    assert len(result.renamed) == 3 and not result.failed and not result.conflicts
    assert (tree / "invoice_acme-invoice.docx").exists()
    assert journal.undo().restored == 3
    assert sorted(x.name for x in tree.rglob("*") if x.name != "log.jsonl") == before


def test_collisions_get_numeric_suffix(tmp_path):
    for n in ("a", "b"):
        make_docx(tmp_path / f"{n}.docx", "Same Title", "meeting agenda", "attendees", "minutes")
    plan = pipeline().plan(tmp_path)
    assert sorted(a.target.name for a in plan.actions) == [
        "meeting-notes_same-title.docx", "meeting-notes_same-title_2.docx"]


def test_rerun_is_idempotent(tree):
    p = pipeline()
    p.apply(p.plan(tree), RenameJournal(tree / 'j.jsonl'))
    assert p.plan(tree).actions == []


def test_date_template_uses_document_metadata(tmp_path):
    make_docx(tmp_path / "a.docx", "Hello")
    plan = pipeline("{date}_{title}").plan(tmp_path)
    assert plan.actions[0].target.name.startswith(str(date.today().year)[:2])


def test_unknown_template_field_rejected():
    with pytest.raises(ValueError):
        NamingConvention("{nope}")


@pytest.mark.parametrize("bad", ["{}", "{0}", "no-fields", "{title}/x", "{title}..{date}", "{nope}"])
def test_invalid_templates_rejected(bad):
    with pytest.raises(ValueError):
        NamingConvention(bad)


def test_hidden_directories_are_skipped(tmp_path):
    (tmp_path / ".git").mkdir()
    make_docx(tmp_path / ".git" / "a.docx", "Invoice", "amount due", "bill to")
    plan = pipeline().plan(tmp_path)
    assert plan.actions == [] and plan.skipped[0].reason is SkipReason.TEMPORARY


def test_journal_survives_failure_midway(tree, monkeypatch):
    from pathlib import Path
    p = pipeline()
    plan = p.plan(tree)
    real = Path.rename
    calls = {"n": 0}

    def flaky(self, target):
        calls["n"] += 1
        if calls["n"] == 2:
            raise PermissionError("locked")
        return real(self, target)

    monkeypatch.setattr(Path, "rename", flaky)
    journal = RenameJournal(tree / "j.jsonl")
    result = p.apply(plan, journal)
    monkeypatch.undo()
    assert len(result.failed) == 1 and len(result.renamed) == 2
    assert journal.undo().restored == 2  # earlier renames are recoverable


def test_apply_reports_conflict_instead_of_overwriting(tree):
    p = pipeline()
    plan = p.plan(tree)
    plan.actions[0].target.write_text("squatter")
    result = p.apply(plan, RenameJournal(tree / "j.jsonl"))
    assert len(result.conflicts) == 1
    assert plan.actions[0].target.read_text() == "squatter"


def test_cli_dry_run_apply_undo(tree, capsys):
    from docclassifier.__main__ import main
    log = str(tree / "log.jsonl")
    assert main(["rename", str(tree)]) == 0
    assert (tree / "scan001.docx").exists()
    assert main(["rename", str(tree), "--apply", "--log", log]) == 0
    assert not (tree / "scan001.docx").exists()
    assert main(["undo", log]) == 0
    assert (tree / "scan001.docx").exists()
    assert main(["rename", str(tree), "--template", "{bad}"]) == 2
    assert main(["undo", str(tree / "missing.jsonl")]) == 2
